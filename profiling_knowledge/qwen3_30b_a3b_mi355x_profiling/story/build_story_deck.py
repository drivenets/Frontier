#!/usr/bin/env python3
"""Build STORY.pptx: the linear_op profiling story as a short deck.

Per iteration: the data we started with -> what did not make sense -> what we found and changed, and how the data
looks now -> what is still open. Figures come from figures/ (make_story_figures.py); captions cite the story's figure numbers.

    ~/.virtualenvs/qwen3-profiling/bin/python build_story_deck.py      # writes STORY.pptx next to this file
"""
import copy
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = lambda n: os.path.join(HERE, "figures", n)

INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x89, 0x87, 0x81)
DARK = RGBColor(0x14, 0x21, 0x3D)
BLUE, ORANGE, AQUA, RED = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xD0, 0x3B, 0x3B)
WHITE, ICE, PALE = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xCA, 0xDC, 0xFC), RGBColor(0x9F, 0xB3, 0xD9)
FONT = "Calibri"
W, H, M = 13.333, 7.5, 0.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
placed = []   # (slide_no, kind, x, y, w, h, text) for the geometry check


def _fill(slide, rgb):
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = rgb


def text(slide, s, x, y, w, h, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    if spacing:
        r.font._element.set("spc", str(spacing))
    placed.append((len(prs.slides), "text", x, y, w, h, s, size))
    return tb


def bullets(slide, items, x, y, w, h, size=14, color=INK, space_after=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.22)))); pPr.set("indent", str(-int(Inches(0.22))))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)
        r = p.add_run(); r.text = it
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
    placed.append((len(prs.slides), "bullets", x, y, w, h, items, size))
    return tb


def image(slide, name, x, y, max_w, max_h):
    with Image.open(FIG(name)) as im:
        ar = im.width / im.height
    w, h = max_w, max_w / ar
    if h > max_h:
        h, w = max_h, max_h * ar
    px, py = x + (max_w - w) / 2, y + (max_h - h) / 2
    slide.shapes.add_picture(FIG(name), Inches(px), Inches(py), Inches(w), Inches(h))
    placed.append((len(prs.slides), "image", px, py, w, h, name, 0))


def kicker(slide, s, color):
    text(slide, s.upper(), M, 0.14, 9, 0.28, size=11, bold=True, color=color, spacing=150)


def title(slide, s, color=INK):
    text(slide, s, M, 0.42, W - 2 * M, 0.7, size=28, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)


def caption(slide, s, x, y, w):
    text(slide, s, x, y, w, 0.42, size=10, color=MUTED)


def footer(slide):
    text(slide, str(len(prs.slides)), W - M - 0.6, H - 0.42, 0.6, 0.28, size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def dark_slide(big, small):
    s = prs.slides.add_slide(BLANK); _fill(s, DARK)
    text(s, big, M, 2.4, W - 2 * M, 0.9, size=40, bold=True, color=WHITE)
    text(s, small, M, 3.35, W - 2 * M, 0.6, size=22, color=ICE)
    return s


def content_slide(k, kcolor, t):
    s = prs.slides.add_slide(BLANK); _fill(s, WHITE)
    kicker(s, k, kcolor); title(s, t)
    return s


# ------------------------------------------------------------------ title
s = prs.slides.add_slide(BLANK); _fill(s, DARK)
text(s, "Qwen3-30B-A3B on MI355X", M, 2.0, W - 2 * M, 0.9, size=40, bold=True, color=WHITE)
text(s, "The linear_op profiling story: three anomalies, three root causes, one clean dataset", M, 2.95, W - 2 * M, 0.7, size=20, color=ICE)
text(s, "Where we started · what did not make sense · what we found and changed · how the data looks now · what is still open", M, 4.2, W - 2 * M, 0.5, size=14, color=PALE)
text(s, "2026-09-15 → 2026-09-23 · Memphis cluster, 8× MI355X · Frontier profilers", M, H - 1.0, W - 2 * M, 0.4, size=12, color=PALE)

# ------------------------------------------------------------------ where we started
s = content_slide("Where we started", BLUE, "The inherited dataset, and our first re-collection")
bullets(s, [
    "Inherited MI355X profiles (Nuzam's data): linear ops up to 4,096 tokens, 20 timed runs per shape, only min/max/mean/median/std kept.",
    "Attention on the portable TORCH_SDPA kernel, no head_dim column; the simulator's regressors train on the per-op median.",
    "Decision: re-collect on the production AITER kernels, 50 timed runs per shape, and keep every run in execution order.",
    "On the shared shapes the medians agree. Everything that follows comes from the raw per-run samples the old data did not have.",
], M, 1.35, 4.6, 5.0, 14)
image(s, "f00_baseline_inherited_vs_50rep.png", 5.4, 1.25, 7.4, 5.2)
caption(s, "Inherited (orange, 20 runs, ≤4,096 tokens) vs our first 50-run collection (blue): attn_pre_proj median per TP. Fig. 1 of the story.", 5.4, 6.55, 7.4)
footer(s)

# ------------------------------------------------------------------ what did not make sense
s = content_slide("The dense grid, 2026-09-15", RED, "Things that did not make sense")
bullets(s, [
    "A single timed run of attn_pre_proj at 150–290 ms, about 1,000× the median, in exactly 32 rows (tokens 1968–1975 × every TP), always timed run 11.",
    "attn_post_proj and attn_rope medians fall 17–37 % between 4,000 and 6,000 tokens at TP 2/4/8 and then climb back: a V in a GEMM cost curve.",
    "Below the V the curves are flat for thousands of tokens although the GEMM's M grows 100×.",
    "Fixed run positions (11, 33, 49) sit above the row median in every row.",
    "None of it is noise: an independent re-run on another node 17 h later reproduced all of it.",
], M, 1.35, 5.0, 5.2, 14)
image(s, "f01_spike_envelope_dense_original.png", 5.8, 1.25, 3.5, 2.6)
image(s, "f09_dip_medians_original.png", 9.4, 1.25, 3.45, 2.6)
image(s, "f04_run_position_fingerprint_original.png", 5.8, 4.0, 7.05, 2.5)
caption(s, "Left: the needle (Fig. 2). Right: the V-shaped dip (Fig. 10). Bottom: the run-position fingerprint (Fig. 4).", 5.8, 6.6, 7.05)
footer(s)

# ================================================================== ITERATION 1
dark_slide("Iteration 1", "The 150–290 ms spike at timed run 11")

s = content_slide("Iteration 1 · the data we started with", BLUE, "One run in fifty, one op, eight token counts, every TP")
image(s, "f01_spike_envelope_dense_original.png", M, 1.25, 6.0, 5.2)
image(s, "f02_spike_one_row_samples.png", 6.8, 1.25, 6.0, 2.5)
bullets(s, [
    "Medians untouched; mean and max destroyed in 32 rows.",
    "Always the same sample position: run 11, the 15th forward including warm-up.",
    "A 26-token repro did not reproduce it; the full grid on another node did, exactly.",
    "So it is keyed to a worker's history, not to the shape or the node.",
], 6.8, 3.95, 6.0, 2.5, 13)
caption(s, "Left: min–max envelope with the spike rows in red (Fig. 2). Right: the 53 samples of one row, run 11 at ≈150 ms (Fig. 3).", M, 6.55, 12.3)
footer(s)

s = content_slide("Iteration 1 · what was the problem", RED, "A CPython full garbage collection inside the timed loop")
bullets(s, [
    "Instrumented every GC collection and the host window of every forward (12 jobs).",
    "Every spiking forward contains one generation-2 collection whose duration equals the event gap to within 1 ms.",
    "Why run 11: the gen-0 counter restarts each task, grows ≈255 + 32 per forward and crosses 700 at forward 14; the full collection fires once per pass at a fixed task index.",
    "Ruled out: caching allocator, hipBLASLt lazy loading, 8-process contention, shape, node.",
    "The event pair measured a host stall while the GPU idled.",
], M, 1.35, 5.2, 5.2, 14)
image(s, "f05_gc_gen2_duration_vs_spike.png", 6.0, 1.25, 6.9, 2.6)
image(s, "f06_gc_counter_model.png", 6.0, 3.95, 6.9, 2.5)
caption(s, "Top: GC duration vs the spiking sample, all 32 on the identity line (Fig. 6). Bottom: the counter arithmetic (Fig. 7).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 1 · what we changed, and how the data looks now", AQUA, "gc.disable() around the timed loop")
bullets(s, [
    "Change: disable CPython's cyclic GC before the warm-up loop, re-enable after the final synchronize (commit 74eee7c).",
    "Verified on the full 13,308-row grid against five pre-stated criteria: no sample above 5 ms, only the sporadic ≤2 ms class left, medians unchanged, fingerprint gone.",
    "Later confirmed with 200 samples per row: the worst single sample in the file is 0.74 ms.",
], M, 1.35, 5.2, 5.2, 14)
image(s, "f07_gcfix_max_over_median_before_after.png", 6.0, 1.25, 6.9, 2.6)
image(s, "f08_fingerprint_before_after_gcfix.png", 6.0, 3.95, 6.9, 2.5)
caption(s, "Top: per-row max ÷ median before and after (Fig. 8). Bottom: the run-11 fingerprint, 1.31 → 1.00 (Fig. 9).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 1 · still open", ORANGE, "The spike is gone. The dip is not.")
bullets(s, [
    "The V in attn_post_proj and attn_rope at TP 2/4/8 is a shift of the row median itself, not a single freak sample; the GC fix cannot touch it.",
    "The plateaus below the dip remain: a GEMM whose M grows 100× should not cost the same.",
    "The bumps at runs 0, 16 and 49 survived the fix: they were never GC.",
], M, 1.35, 5.2, 5.0, 14)
image(s, "f09_dip_medians_original.png", 6.0, 1.35, 6.9, 4.9)
caption(s, "Medians vs tokens, three ops × four TPs, dip band shaded (Fig. 10).", 6.0, 6.4, 6.9)
footer(s)

# ================================================================== ITERATION 2
dark_slide("Iteration 2", "The V-shaped dip and the flat plateaus")

s = content_slide("Iteration 2 · the data we started with", BLUE, "A cost that falls when the work grows")
image(s, "f09_dip_medians_original.png", M, 1.25, 7.2, 2.8)
image(s, "f10_dip_row_samples_bimodal.png", M, 4.05, 7.2, 2.4)
bullets(s, [
    "attn_post_proj drops 35–37 % at TP 2 around 4,100 tokens; attn_rope 25–27 % at TP 4 near 5,900. TP 1 never dips.",
    "First hypothesis: hipBLASLt picks a faster tile for the K = 4096/TP GEMMs. Both leading candidates assumed the number in the CSV is GPU time.",
    "The clue: inside the dip a row's 50 samples are bimodal and drift downward within the task; past the dip they are tight.",
], 8.0, 1.35, 4.85, 5.2, 13)
caption(s, "Top: the dip (Fig. 10). Bottom: samples of a dip row vs a clean row, traced kernel dashed (Fig. 11).", M, 6.55, 7.2)
footer(s)

s = content_slide("Iteration 2 · what was the problem", RED, "The profiler was measuring the host, not the GPU")
bullets(s, [
    "A CUDA-event pair times from when the device reaches the start event to when it reaches the end event. If the device is idle when the start event is enqueued, the pair measures the host's launch span.",
    "One knob, no dip: holding the device 100 ms behind the host made every curve monotonic and equal to the traced kernel + ≈6 µs.",
    "Kernel traces: the o_proj kernel is monotonic through the dip; the GPU was idle 45–60 % of every forward at TP 2/4/8.",
    "Not confined to the dip: the old medians were 1.2–3× GPU time over most of the TP>1 grid.",
    "Validated against a pre-registered prediction with seven falsification conditions before any fix was applied.",
], M, 1.35, 5.2, 5.3, 13)
image(s, "f11_dip_original_vs_gpu_backlog.png", 6.0, 1.25, 6.9, 2.8)
image(s, "f12_inflation_original_over_gpu_bound.png", 6.0, 4.15, 6.9, 2.3)
caption(s, "Top: original vs device-held-busy medians (Fig. 12). Bottom: how inflated, by op × TP × tokens (Fig. 13).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 2 · what we changed, and how the data looks now", AQUA, "Two timing columns, a clock probe, the real RoPE kernel")
bullets(s, [
    "Every shape timed twice: the legacy loop (kept as a record) and a GPU-bound pass behind a device backlog, re-armed every 25 forwards.",
    "A shader-clock probe on every row: warm-up did not warm the clock; the legacy loop started at ≈0.8 GHz.",
    "attn_rope had timed a numerically wrong torch fallback; replaced by the fused kernel, 2–8× cheaper.",
    "Measurement contract: the label is kernel time at a stated clock plus a separate launch-overhead term.",
    "Re-collected dense grid: no dip, no plateaus; the 15–30 % steps that remain are real hipBLASLt tile switches.",
], M, 1.35, 5.2, 5.3, 13)
image(s, "f18_dip_gone_old_vs_fixed.png", 6.0, 1.25, 6.9, 3.2)
image(s, "f16_rope_fallback_vs_fused.png", 6.0, 4.5, 6.9, 1.95)
caption(s, "Top: old canonical (orange) vs fixed GPU-bound (blue): the dip is gone (Fig. 19). Bottom: RoPE fallback vs fused kernel (Fig. 17).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 2 · still open", ORANGE, "Clean medians, but the per-run samples show two new patterns")
bullets(s, [
    "Within every 25-run block the samples decline 2–3 % from run 3 to run 25; the last run is the minimum far more often than chance.",
    "At TP>1 attn_pre_proj has its maximum at run 20 in 42–60 % of rows, and at run 118 in the 200-run file.",
    "Medians are barely affected, but min/max/mean/std are, and a settled number should not depend on where in the block it was taken.",
], M, 1.35, 5.2, 5.0, 14)
image(s, "f23_position_profile_fixed_25.png", 6.0, 1.35, 6.9, 4.9)
caption(s, "Run-position profile of attn_pre_proj in the fixed collection: GPU-bound column left, legacy column right (Fig. 24).", 6.0, 6.4, 6.9)
footer(s)

# ================================================================== ITERATION 3
dark_slide("Iteration 3", "Run-position patterns in the corrected data")

s = content_slide("Iteration 3 · the data we started with", BLUE, "A ramp in every block, a spike at a fixed position")
image(s, "f24_position_profile_200_reps.png", M, 1.25, 7.6, 5.3)
bullets(s, [
    "200 runs as 8 spun blocks of 25: the downward ramp restarts after every re-armed spin.",
    "Spikes fall exactly at runs 20, 40, 59, 79, 98, 118, 138, 157, 177, 196: predicted in advance by one barrier per 1000 commands since the last synchronize.",
    "Two separate mechanisms, both invisible in medians.",
], 8.4, 1.35, 4.45, 5.2, 13)
caption(s, "attn_pre_proj and whole-forward span over 200 runs; dotted = spin re-armed, red = predicted barrier positions (Fig. 25).", M, 6.55, 7.6)
footer(s)

s = content_slide("Iteration 3 · what was the problem", RED, "The idle spin drops the clocks; the runtime adds a barrier")
bullets(s, [
    "Effect 1, the ramp: the torch.cuda._sleep spin used as backlog is a low-activity state; the clock domains idle and the first 10–15 forwards of every block run while they ramp back up. Replacing the spin by real GEMMs removes it in single-GPU probes.",
    "Effect 2, the spike: ROCclr enqueues a clean-up marker after every 1000 commands; on the device it is a barrier with system-scope fences costing ≈5.5 µs, landing in forward 20's attn_pre_proj at TP>1. Confirmed in a kernel trace, the runtime log and the ROCclr source.",
], M, 1.35, 5.2, 5.3, 13)
image(s, "f26_clock_ramp_after_spin.png", 6.0, 1.25, 6.9, 2.5)
image(s, "f29_effect1_fix_variants.png", 6.0, 3.85, 6.9, 2.6)
caption(s, "Top: clock after the spin, and spin vs GEMM backlog (Fig. 26). Bottom: the fix campaign; only a work backlog removes the drift (Fig. 28).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 3 · what we changed, and how the data looks now", AQUA, "Work backlog + 8 settle forwards + a runtime flag")
bullets(s, [
    "DEBUG_CLR_MAX_BATCH_SIZE=1000000 in the container: the barrier never lands inside a timed block; the run-20 spike and the legacy stall are gone.",
    "Backlog = an allocating chain of 4096² GEMMs (settle-only after the spin was ruled out; two other backlog kinds tried), then 8 untimed settle forwards.",
    "Dense grid: profiles flat from run 1 at TP 4/8; medians 2–3 % lower than the spin collections, the settled clock; per-row spread 40 % smaller.",
    "This file (job 21539) is the one handed to the regressor, with a checksum and a handoff document defining the label.",
], M, 1.35, 5.2, 5.3, 13)
image(s, "f27_settle_position_profiles.png", 6.0, 1.25, 6.9, 2.1)
image(s, "f30_final_dataset_medians.png", 6.0, 3.45, 6.9, 3.0)
caption(s, "Top: spin vs work backlog vs +8 settle, per TP (Fig. 29). Bottom: the handed-off dataset (Fig. 31).", 6.0, 6.55, 6.9)
footer(s)

s = content_slide("Iteration 3 · still open", ORANGE, "What we still have")
bullets(s, [
    "A slow 0.9–1.3 % decline over the block at TP 1/2 in 8-GPU runs, independent of the settle count and absent on a single GPU: probably clock or thermal settling under full-node load. Medians move < 0.5 %.",
    "Kernels below ≈10 µs sit inside the instrument's ≈3 µs floor and are not cross-validated between event pairs and traces.",
    "The locked-clock dataset the contract calls for has not been collected.",
    "The dataset README names the 200-forward file as canonical while the handoff points at the settle-8 file; one has to give.",
    "The attention datasets never had this per-run scrutiny.",
], M, 1.35, 6.4, 5.3, 14)
image(s, "f28_settled_vs_spin_medians.png", 7.2, 1.35, 5.7, 4.6)
caption(s, "Settled ÷ spin medians: what the ramp had cost (Fig. 30).", 7.2, 6.1, 5.7)
footer(s)

# ------------------------------------------------------------------ closing
s = prs.slides.add_slide(BLANK); _fill(s, DARK)
text(s, "What we learned", M, 0.6, W - 2 * M, 0.8, size=32, bold=True, color=WHITE)
bullets(s, [
    "Keep every sample. All three anomalies were invisible in medians and obvious in per-run data.",
    "Write the prediction down before the run. The dip conclusion, the barrier positions and the fix acceptance criteria were all committed in advance.",
    "State what the number is. The measurement contract, the clock probe and the two columns exist so a per-op sum never mixes host spans with kernels.",
    "The old file overstated GEMMs 1.3–2.9× and RoPE 4–12×. The new one is monotonic, instrument-accounted and 2–3 % lower again once the clock was settled.",
], M, 1.7, W - 2 * M, 4.2, 18, color=WHITE, space_after=12)
text(s, "Full story, figures and sources: profiling_knowledge/qwen3_30b_a3b_mi355x_profiling/story/STORY.md · Confluence page 7166984214", M, H - 1.0, W - 2 * M, 0.4, size=12, color=PALE)

out = os.path.join(HERE, "STORY.pptx")
prs.save(out)
print(f"wrote {out}: {len(prs.slides)} slides")

# ------------------------------------------------------------------ geometry check (no renderer on this machine)
# 1. nothing outside the slide; 2. no overlapping boxes on a slide; 3. rough text-fit estimate (Calibri ≈ 0.48 em average width).
problems = []
by_slide = {}
for rec in placed:
    by_slide.setdefault(rec[0], []).append(rec)
for sn, items in by_slide.items():
    for (_, kind, x, y, w, h, payload, size) in items:
        if x < 0.3 or y < 0.05 or x + w > W - 0.3 + 1e-6 or y + h > H - 0.1 + 1e-6:
            problems.append(f"slide {sn}: {kind} outside margins ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
        if kind in ("text", "bullets"):
            lines = payload if isinstance(payload, list) else [payload]
            chars_per_line = max(1, int(w * 72 / (size * 0.48)))
            n_lines = sum(-(-len(t) // chars_per_line) for t in lines)
            need = n_lines * size * 1.2 / 72 + (len(lines) - 1) * 6 / 72
            if need > h * 1.02:
                problems.append(f"slide {sn}: text may overflow: needs ≈{need:.2f} in, box {h:.2f} in — {str(lines[0])[:50]}…")
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            a, b = items[i], items[j]
            ax, ay, aw, ah = a[2:6]; bx, by, bw, bh = b[2:6]
            if ax < bx + bw - 0.02 and bx < ax + aw - 0.02 and ay < by + bh - 0.02 and by < ay + ah - 0.02:
                problems.append(f"slide {sn}: overlap {a[1]}({str(a[6])[:25]}) × {b[1]}({str(b[6])[:25]})")
print("\n".join(problems) if problems else "geometry check: no overlaps, nothing outside margins, text fits (estimate)")
